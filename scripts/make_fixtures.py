"""Generate the committed test fixtures in tests/fixtures/.

Run once; the outputs are committed so tests stay deterministic:

    uv run python scripts/make_fixtures.py
"""

from __future__ import annotations

import base64
import sys
import zipfile
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

FIXTURES = Path(__file__).parent.parent / "tests" / "fixtures"

# A valid 1x1 red PNG (vision extraction is faked in tests, so content is irrelevant)
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGP4z8DwHwAFAAH/q842iQAAAABJRU5ErkJggg=="
)


def make_sample_pdf() -> None:
    """Two pages with real extractable text."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    c.drawString(72, 720, "Photosynthesis converts light energy into chemical energy.")
    c.drawString(72, 700, "Chlorophyll absorbs blue and red light most efficiently.")
    c.showPage()
    c.drawString(72, 720, "The Calvin cycle fixes carbon dioxide into glucose.")
    c.showPage()
    c.save()
    (FIXTURES / "sample.pdf").write_bytes(buf.getvalue())


def make_scanned_pdf() -> None:
    """Two pages with NO extractable text (image-only / blank) -> triggers OCR."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=letter)
    # Draw rectangles only - no text layer, like a scan without OCR
    c.rect(72, 600, 400, 100, fill=1)
    c.showPage()
    c.rect(72, 600, 400, 100, fill=1)
    c.showPage()
    c.save()
    (FIXTURES / "scanned.pdf").write_bytes(buf.getvalue())


def make_sample_docx() -> None:
    import docx

    document = docx.Document()
    document.add_heading("Mitochondria", level=1)
    document.add_paragraph("Mitochondria are the powerhouse of the cell.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Organelle"
    table.rows[0].cells[1].text = "Function"
    table.rows[1].cells[0].text = "Ribosome"
    table.rows[1].cells[1].text = "Protein synthesis"
    document.save(FIXTURES / "sample.docx")


def make_sample_xlsx() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Grades"
    ws.append(["Student", "Subject", "Grade"])
    ws.append(["Ana", "Math", 95])
    ws.append(["Luis", "Biology", 88])
    wb.save(FIXTURES / "sample.xlsx")


def main() -> None:
    FIXTURES.mkdir(parents=True, exist_ok=True)
    make_sample_pdf()
    make_scanned_pdf()
    make_sample_docx()
    make_sample_xlsx()
    (FIXTURES / "sample.png").write_bytes(PNG_1X1)
    (FIXTURES / "sample.txt").write_text(
        "Newton's first law states that objects in motion stay in motion."
    )
    # Sanity: every fixture exists and is non-empty
    for name in [
        "sample.pdf",
        "scanned.pdf",
        "sample.docx",
        "sample.xlsx",
        "sample.png",
        "sample.txt",
    ]:
        path = FIXTURES / name
        assert path.stat().st_size > 0, name
    # docx/xlsx are zips - verify they open
    for name in ["sample.docx", "sample.xlsx"]:
        assert zipfile.ZipFile(FIXTURES / name).namelist()
    print(f"Fixtures written to {FIXTURES}")


if __name__ == "__main__":
    main()
