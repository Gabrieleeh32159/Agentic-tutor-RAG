from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.ingestion.base import ParsedBlock, ParsedDocument
from app.ingestion.office import check_zip_bomb
from app.shared.errors import ParseFailedError

MAX_SHEETS = 10
MAX_ROWS_PER_SHEET = 2000
ROWS_PER_BLOCK = 30


def _markdown_row(values: tuple) -> str:
    cells = [
        "" if v is None else str(v).replace("\n", " ").replace("|", "\\|")
        for v in values
    ]
    return "| " + " | ".join(cells) + " |"


class XlsxParser:
    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        check_zip_bomb(data)
        try:
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ParseFailedError("Could not read the Excel file") from exc

        blocks: list[ParsedBlock] = []
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                continue
            header_md = _markdown_row(header)
            separator = "|" + "---|" * len(header)

            batch: list[str] = []
            row_count = 0
            for values in rows:
                if row_count >= MAX_ROWS_PER_SHEET:
                    break
                if all(v is None for v in values):
                    continue
                batch.append(_markdown_row(values))
                row_count += 1
                if len(batch) == ROWS_PER_BLOCK:
                    blocks.append(
                        ParsedBlock(
                            text="\n".join([header_md, separator, *batch]),
                            sheet_name=sheet.title,
                            is_atomic=True,
                        )
                    )
                    batch = []
            if batch:
                blocks.append(
                    ParsedBlock(
                        text="\n".join([header_md, separator, *batch]),
                        sheet_name=sheet.title,
                        is_atomic=True,
                    )
                )
        workbook.close()

        if not blocks:
            raise ParseFailedError("The Excel file contains no data")

        return ParsedDocument(blocks=blocks)
