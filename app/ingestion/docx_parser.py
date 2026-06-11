from __future__ import annotations

from io import BytesIO

import docx

from app.ingestion.base import ParsedBlock, ParsedDocument
from app.ingestion.office import check_zip_bomb
from app.shared.errors import ParseFailedError


def _table_to_markdown(table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    lines = ["| " + " | ".join(rows[0]) + " |"]
    lines.append("|" + "---|" * len(rows[0]))
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


class DocxParser:
    def parse(self, data: bytes, filename: str) -> ParsedDocument:
        check_zip_bomb(data)
        try:
            document = docx.Document(BytesIO(data))
        except Exception as exc:
            raise ParseFailedError("Could not read the Word document") from exc

        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text.strip())
        for table in document.tables:
            markdown = _table_to_markdown(table)
            if markdown:
                parts.append(markdown)

        if not parts:
            raise ParseFailedError("The Word document contains no text")

        return ParsedDocument(blocks=[ParsedBlock(text="\n\n".join(parts))])
